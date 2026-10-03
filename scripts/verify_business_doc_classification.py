"""
Verification Script for Business Document Classification Bug Fix.
Tests end-to-end classification of C:\\Users\\srivi\\Downloads\\Business_Test_Article.docx
and checks cross-category separation across all 5 canonical categories.
"""
import os
import sys
import docx
import requests
import json

DOCX_PATH = r"C:\Users\srivi\Downloads\Business_Test_Article.docx"
BASE_URL = "http://127.0.0.1:8000/api"

def main():
    print("=" * 70)
    print("RUNNING END-TO-END BUSINESS DOCUMENT CLASSIFICATION VERIFICATION")
    print("=" * 70)
    
    # 1. Verify file exists
    assert os.path.exists(DOCX_PATH), f"Target docx file not found: {DOCX_PATH}"
    doc = docx.Document(DOCX_PATH)
    full_text = "\n\n".join([p.text for p in doc.paragraphs if p.text.strip()])
    print(f"\n[1] File verified: {DOCX_PATH}")
    print(f"    Total paragraphs: {len(doc.paragraphs)}")
    print(f"    Total text length: {len(full_text)} characters")
    print(f"    Headline: {repr(doc.paragraphs[0].text)}")

    # 2. Test POST /api/extract (State 2 Preview)
    with open(DOCX_PATH, "rb") as f:
        resp = requests.post(f"{BASE_URL}/extract", files={"file": ("Business_Test_Article.docx", f, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
    assert resp.status_code == 200, f"/api/extract failed: {resp.text}"
    extract_data = resp.json()
    print(f"\n[2] /api/extract response:")
    print(f"    Extracted headline: {extract_data.get('headline')}")
    print(f"    Char count: {extract_data.get('char_count')}")
    print(f"    Word count: {extract_data.get('word_count')}")
    assert "economic" in extract_data["text"].lower()

    # 3. Test POST /api/predict/file (Document Classification)
    with open(DOCX_PATH, "rb") as f:
        resp = requests.post(
            f"{BASE_URL}/predict/file",
            files={"file": ("Business_Test_Article.docx", f, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
            data={"model_name": "best"}
        )
    assert resp.status_code == 200, f"/api/predict/file failed: {resp.text}"
    result = resp.json()
    
    print(f"\n[3] /api/predict/file result for Business_Test_Article.docx:")
    print(f"    Model Used:         {result.get('model_used')}")
    print(f"    Predicted Category: {result.get('predicted_category')}")
    print(f"    Confidence:         {result.get('confidence') * 100:.2f}%")
    print(f"    All Probabilities:  {json.dumps(result.get('probabilities'), indent=6)}")
    
    # Assertions for Business Document
    predicted = result.get("predicted_category")
    probs = result.get("probabilities", {})
    bus_prob = probs.get("Business", 0.0)
    tech_prob = probs.get("Technology", 0.0)
    
    assert predicted == "Business", f"FAILURE: Expected Business, got {predicted}"
    assert bus_prob > tech_prob, f"FAILURE: Business prob ({bus_prob}) should be > Tech prob ({tech_prob})"
    assert bus_prob > 0.50, f"FAILURE: Business prob ({bus_prob}) should be > 50%"
    print("\n>>> SUCCESS: Business_Test_Article.docx correctly classified as BUSINESS with > 50% calibrated confidence! <<<")

    # 4. Test other models individually
    for m in ["Naive Bayes", "Logistic Regression", "Support Vector Machine"]:
        with open(DOCX_PATH, "rb") as f:
            resp = requests.post(
                f"{BASE_URL}/predict/file",
                files={"file": ("Business_Test_Article.docx", f, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
                data={"model_name": m}
            )
        res = resp.json()
        print(f"    Model [{m}]: Predicted={res['predicted_category']} (Conf={res['confidence']*100:.1f}%, Bus={res['probabilities']['Business']}, Tech={res['probabilities']['Technology']})")
        assert res["predicted_category"] == "Business", f"Model {m} failed to predict Business!"

    # 5. Cross-Category Separation Verification
    print("\n[5] Testing Cross-Category Separation across all canonical categories:")
    test_cases = [
        {
            "cat": "Technology",
            "headline": "Nvidia Announces Blackwell Ultra High Performance AI Superchip",
            "text": "Nvidia unveiled its newest Blackwell GPU supercomputing architecture featuring fifth-generation tensor cores, ultra-low latency NVLink interconnects, and dedicated fp4 acceleration designed specifically for deep neural network training and large scale generative AI model inference."
        },
        {
            "cat": "Education",
            "headline": "Universities Introduce New STEM Degree Curriculum and Scholarship Standards",
            "text": "Colleges and universities nationwide are overhauling undergraduate STEM curricula with expanded accreditation requirements, student faculty mentorship programs, and academic scholarships to boost retention in computer science and mathematics."
        },
        {
            "cat": "Entertainment",
            "headline": "Hollywood Studio Breaks Box Office Record Ahead of Oscar Nominations",
            "text": "The highly anticipated cinematic release scored a massive opening weekend at the worldwide box office, earning glowing film reviews from critics and setting new records for theatrical attendance as awards season approaches."
        },
        {
            "cat": "Sports",
            "headline": "Championship Finals Conclude with Historic Double Overtime Victory",
            "text": "In a thrilling championship game that went into double overtime, the visiting team clinched the league title with a buzzer-beater goal as thousands of cheering fans celebrated in the stadium."
        }
    ]
    
    for tc in test_cases:
        p_resp = requests.post(
            f"{BASE_URL}/predict",
            json={"text": tc["text"], "headline": tc["headline"], "model_name": "best"}
        )
        assert p_resp.status_code == 200
        p_res = p_resp.json()
        pred_cat = p_res["predicted_category"]
        conf = p_res["confidence"]
        print(f"    Expected: {tc['cat']:<15} | Predicted: {pred_cat:<15} | Confidence: {conf*100:.2f}%")
        assert pred_cat == tc["cat"], f"Expected {tc['cat']}, got {pred_cat}"
        if tc["cat"] == "Technology":
            assert conf > 0.80, f"Technology confidence ({conf}) should be > 80%"

    print("\n==================================================================")
    print("ALL VERIFICATION CHECKS PASSED: 100% ACCURACY & CATEGORY INTEGRITY!")
    print("==================================================================")

if __name__ == "__main__":
    main()
